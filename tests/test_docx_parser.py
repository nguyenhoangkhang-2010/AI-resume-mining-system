import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor
from lxml import etree

from app.evaluation.cv_corpus import SECTION_TITLES, CVCorpusBuilder, TruthLine
from app.evaluation.section_eval import aggregate, score_document
from app.extraction.rules.experience_rules import parse_experience
from app.extraction.section.section_detector import SectionDetector
from app.resume_processing.layout.layout_analyzer import LayoutAnalyzer
from app.resume_processing.layout.layout_representation import LayoutDocument
from app.resume_processing.layout.semantic_region_analyzer import SemanticRegionAnalyzer
from app.resume_processing.parsers.docx_parser import DocxParser, UnreadableFileError
from app.resume_processing.parsers.document_parser import DocumentParser
from app.resume_processing.pipelines.resume_pipeline import ResumePipeline
from app.services.resume_quality import UNREADABLE_FILE, failure_reason


def lines_of(parsed):
    return [
        line
        for page in parsed["pages"]
        for block in page["blocks"]
        for line in block["lines"]
    ]


def save(document, tmp_path, name="cv.docx"):
    path = tmp_path / name
    document.save(str(path))

    return path


@pytest.fixture
def simple(tmp_path):
    document = Document()
    document.add_heading("Nguyen Van A", 0)
    document.add_paragraph("a@b.co | +84 912 345 678 | Hanoi")
    document.add_heading("Experience", 1)
    title = document.add_paragraph()
    title.add_run("Backend Engineer - Acme Systems").bold = True
    document.add_paragraph("2022 - 2025")
    document.add_paragraph("Built payment services", style="List Bullet")
    document.add_paragraph("Designed APIs", style="List Bullet")
    document.add_heading("Skills", 1)
    document.add_paragraph("Python, SQL, Docker")

    return save(document, tmp_path)


# ------------------------------------------------------------------ reading


def test_text_comes_out_in_document_order_with_visible_bullets(simple):
    text = DocxParser.parse(simple)["text"].splitlines()

    assert text == [
        "Nguyen Van A",
        "a@b.co | +84 912 345 678 | Hanoi",
        "Experience",
        "Backend Engineer - Acme Systems",
        "2022 - 2025",
        "• Built payment services",
        "• Designed APIs",
        "Skills",
        "Python, SQL, Docker",
    ]


def test_the_result_has_the_shape_the_layout_stage_expects(simple):
    parsed = DocxParser.parse(simple)

    assert set(parsed) == {"text", "pages", "metadata"}
    page = parsed["pages"][0]
    assert (page["width"], page["height"]) == (612.0, 792.0)  # python-docx default: Letter
    assert set(page["blocks"][0]) == {"type", "block_index", "bbox", "text", "lines"}
    assert set(page["blocks"][0]["lines"][0]) == {"text", "bbox", "line_index", "spans"}
    assert {"size", "flags", "font", "color", "bbox"} <= set(
        page["blocks"][0]["lines"][0]["spans"][0]
    )
    assert LayoutAnalyzer().analyze(LayoutDocument.from_parser_result(parsed)).block_count == 9


def test_headings_and_bold_runs_are_marked_bold_and_larger(simple):
    by_text = {line["text"]: line for line in lines_of(DocxParser.parse(simple))}

    heading = by_text["Experience"]["spans"][0]
    body = by_text["Python, SQL, Docker"]["spans"][0]
    bold_run = by_text["Backend Engineer - Acme Systems"]["spans"][0]

    assert heading["flags"] & 16 and bold_run["flags"] & 16
    assert not body["flags"] & 16
    assert heading["size"] > body["size"]


def test_explicit_sizes_colours_and_italics_are_kept(tmp_path):
    document = Document()
    run = document.add_paragraph().add_run("Styled")
    run.font.size = Pt(18)
    run.italic = True
    run.font.color.rgb = RGBColor(0x10, 0x20, 0x30)

    span = lines_of(DocxParser.parse(save(document, tmp_path)))[0]["spans"][0]

    assert span["size"] == 18.0
    assert span["flags"] & 2
    assert span["color"] == 0x102030


def test_space_before_a_heading_becomes_a_gap_between_lines(tmp_path):
    document = Document()
    document.add_paragraph("above")
    heading = document.add_paragraph("Skills")
    heading.paragraph_format.space_before = Pt(24)
    document.add_paragraph("below")

    lines = {line["text"]: line for line in lines_of(DocxParser.parse(save(document, tmp_path)))}

    gap = lines["Skills"]["bbox"]["y0"] - lines["above"]["bbox"]["y1"]
    tight = lines["below"]["bbox"]["y0"] - lines["Skills"]["bbox"]["y1"]

    assert gap > 20 and tight < gap


def test_a_line_break_inside_a_paragraph_makes_two_lines(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("first")
    paragraph.add_run().add_break()
    paragraph.add_run("second")

    texts = [line["text"] for line in lines_of(DocxParser.parse(save(document, tmp_path)))]

    assert texts == ["first", "second"]


def test_a_link_keeps_its_text(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Mail ")
    body = paragraph._p
    link = etree.SubElement(
        body,
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink",
    )
    run = etree.SubElement(
        link, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"
    )
    text = etree.SubElement(
        run, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    )
    text.text = "a@b.co"

    assert DocxParser.parse(save(document, tmp_path))["text"] == "Mail a@b.co"


def test_page_breaks_start_new_pages(tmp_path):
    document = Document()
    document.add_paragraph("page one")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("page two")

    parsed = DocxParser.parse(save(document, tmp_path))

    assert [page["text"] for page in parsed["pages"]] == ["page one", "page two"]


def test_a_long_document_runs_onto_more_pages(tmp_path):
    document = Document()

    for number in range(120):
        document.add_paragraph(f"Line number {number} of a long document")

    assert len(DocxParser.parse(save(document, tmp_path))["pages"]) >= 2


def test_header_text_is_read_and_comes_first(tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Nguyen Van A  a@b.co"
    document.add_paragraph("Body text")

    assert DocxParser.parse(save(document, tmp_path))["text"].splitlines() == [
        "Nguyen Van A  a@b.co",
        "Body text",
    ]


def test_a_document_with_no_text_gives_empty_text(tmp_path):
    parsed = DocxParser.parse(save(Document(), tmp_path))

    assert parsed["text"] == ""


# ------------------------------------------------------------ tables, boxes


def two_column_table(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.paragraphs[0].text = "Skills"
    left.add_paragraph("Python, SQL")
    left.add_paragraph("Certificates")
    left.add_paragraph("AWS Practitioner")
    right.paragraphs[0].text = "Experience"
    right.add_paragraph("Backend Engineer - Acme Systems")
    right.add_paragraph("2022 - 2025")
    right.add_paragraph("Education")
    right.add_paragraph("BSc in Computer Science")

    return save(document, tmp_path)


def test_a_two_column_table_lays_out_as_two_columns(tmp_path):
    parsed = DocxParser.parse(two_column_table(tmp_path))
    analysis = LayoutAnalyzer().analyze(LayoutDocument.from_parser_result(parsed))
    columns = SemanticRegionAnalyzer()._detect_columns(
        analysis.pages[0].blocks, analysis.pages[0].width
    )

    assert len(columns) == 2

    left, right = ({b.text.split("\n")[0] for b in column} for column in columns)

    assert "Skills" in left and "Experience" in right


def test_a_table_below_text_continues_down_the_page(tmp_path):
    document = Document()
    document.add_paragraph("Top line")
    table = document.add_table(rows=2, cols=1)
    table.rows[0].cells[0].text = "row one"
    table.rows[1].cells[0].text = "row two"

    ys = {line["text"]: line["bbox"]["y0"] for line in lines_of(DocxParser.parse(save(document, tmp_path)))}

    assert ys["Top line"] < ys["row one"] < ys["row two"]


TEXT_BOX = """
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:r>
    <mc:AlternateContent>
      <mc:Choice>
        <w:drawing>
          <wp:anchor>
            <wp:positionH relativeFrom="page"><wp:posOffset>381000</wp:posOffset></wp:positionH>
            <wp:positionV relativeFrom="page"><wp:posOffset>1270000</wp:posOffset></wp:positionV>
            <wp:extent cx="2032000" cy="3000000"/>
            <a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
              <a:graphicData><wps:wsp><wps:txbx><w:txbxContent>
                <w:p><w:r><w:t>Languages</w:t></w:r></w:p>
                <w:p><w:r><w:t>English, Vietnamese</w:t></w:r></w:p>
              </w:txbxContent></wps:txbx></wps:wsp></a:graphicData>
            </a:graphic>
          </wp:anchor>
        </w:drawing>
      </mc:Choice>
      <mc:Fallback>
        <w:pict><w:txbxContent>
          <w:p><w:r><w:t>Languages</w:t></w:r></w:p>
          <w:p><w:r><w:t>English, Vietnamese</w:t></w:r></w:p>
        </w:txbxContent></w:pict>
      </mc:Fallback>
    </mc:AlternateContent>
  </w:r>
</w:p>
"""


def test_text_in_a_box_is_read_once_and_placed_where_it_is_anchored(tmp_path):
    document = Document()
    document.add_paragraph("Body text")
    document.element.body.insert(0, etree.fromstring(TEXT_BOX))

    parsed = DocxParser.parse(save(document, tmp_path))
    found = {line["text"]: line for line in lines_of(parsed)}

    assert parsed["text"].count("Languages") == 1
    assert found["English, Vietnamese"]["bbox"]["x0"] == pytest.approx(30.0, abs=0.5)
    assert found["Languages"]["bbox"]["y0"] == pytest.approx(100.0, abs=0.5)


# ------------------------------------------------------------------ unsafe


def test_a_missing_file_is_reported(tmp_path):
    with pytest.raises(FileNotFoundError):
        DocxParser.parse(tmp_path / "absent.docx")


@pytest.mark.parametrize("content", [b"", b"PK\x03\x04not really a zip", b"%PDF-1.4 hello"])
def test_a_file_that_is_not_a_word_document_is_refused(tmp_path, content):
    path = tmp_path / "bad.docx"
    path.write_bytes(content)

    with pytest.raises(UnreadableFileError):
        DocxParser.parse(path)


def test_a_zip_without_a_word_document_is_refused(tmp_path):
    path = tmp_path / "other.docx"

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "hello")

    with pytest.raises(UnreadableFileError, match="not a Word document"):
        DocxParser.parse(path)


def test_an_archive_that_expands_enormously_is_refused(tmp_path):
    path = tmp_path / "bomb.docx"

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "<w:document/>")
        archive.writestr("padding.bin", b"0" * (60 * 1024 * 1024))

    with pytest.raises(UnreadableFileError, match="too large"):
        DocxParser.parse(path)


def test_a_broken_word_file_is_reported_as_unreadable(tmp_path):
    path = tmp_path / "broken.docx"

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "<<<not xml")

    with pytest.raises(UnreadableFileError):
        DocxParser.parse(path)


def test_the_failure_reason_for_an_unreadable_word_file_is_stored_as_unreadable():
    assert failure_reason(UnreadableFileError("secret content")) == UNREADABLE_FILE


# ------------------------------------------------------------- the dispatcher


def test_the_parser_follows_the_file_type(simple, make_pdf):
    pdf = make_pdf(["Hello from a PDF"])

    assert "Hello from a PDF" in DocumentParser().parse(pdf)["text"]
    assert "Backend Engineer" in DocumentParser().parse(simple)["text"]


def test_other_types_are_refused(tmp_path):
    path = tmp_path / "cv.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0")

    with pytest.raises(UnreadableFileError, match="Unsupported"):
        DocumentParser().parse(path)


# ------------------------------------------------------ the whole pipeline


class NeverConfident:
    def classify(self, text):
        return "other", 0.0


class NoLLM:
    def classify_single(self, text):
        return "other", 0.0

    def classify_document(self, lines):
        return []


class EchoEngine:
    def extract_resume(self, sections):
        return {"section_names": sorted(k for k in sections if not k.startswith("__"))}


def pipeline():
    return ResumePipeline(
        SectionDetector(NeverConfident(), NoLLM()),
        extraction_engine=EchoEngine(),  # type: ignore[arg-type]
    )


def test_a_word_cv_goes_through_the_whole_pipeline(simple):
    result = pipeline().process_pdf(simple)

    assert "Backend Engineer" in result["raw_text"]
    assert result["section_names"] == ["experience", "personal_info", "skills"]


def test_the_rules_read_the_jobs_of_a_word_cv(simple):
    sections = pipeline().detect_sections(simple)

    jobs = parse_experience("\n".join(sections["experience"]))

    assert [(job["role"], job["company"]) for job in jobs] == [("Backend Engineer", "Acme Systems")]
    assert jobs[0]["description"] == ["Built payment services", "Designed APIs"]


def docx_from_truth(document, directory: Path, language: str = "en"):
    """
    A Word version of a corpus CV: the same lines under Heading 1 titles.
    Returns the file and the corpus document with its truth in the form the
    parser gives back (a list bullet comes out as "• ").
    """

    word = Document()
    sections: dict = {}

    for line in document.truth:
        sections.setdefault(line.section, []).append(line.text)

    for text in sections["personal_info"][:1]:
        word.add_heading(text, 0)

    for text in sections["personal_info"][1:]:
        word.add_paragraph(text)

    for section in ("summary", "experience", "education", "skills", "projects", "certifications"):
        word.add_heading(SECTION_TITLES[language][section], 1)

        for text in sections[section]:
            if text.startswith("- "):
                word.add_paragraph(text[2:], style="List Bullet")
            else:
                word.add_paragraph(text)

    path = directory / f"{document.id}.docx"
    word.save(str(path))

    shown = replace(
        document,
        truth=[
            TruthLine(line.section, "• " + line.text[2:] if line.text.startswith("- ") else line.text)
            for line in document.truth
        ],
    )

    return path, shown


def test_sections_are_found_in_generated_word_cvs(tmp_path):
    corpus = CVCorpusBuilder(seed=3, per_variant=1).build(tmp_path / "pdf")
    documents = [d for d in corpus if d.wording == "standard" and d.language == "en"][:6]
    results = []

    for document in documents:
        docx, shown = docx_from_truth(document, tmp_path)
        results.append(score_document(shown, pipeline().detect_sections(docx)))

    assert aggregate(results)["all"].f1 >= 0.95


# ----------------------------------------------------- upload and the API


def word_bytes(tmp_path):
    document = Document()
    document.add_paragraph("Nguyen Van A")
    path = save(document, tmp_path, "upload.docx")

    return path.read_bytes()


def test_a_real_word_file_is_accepted_on_upload(tmp_path):
    import io

    from app.services.upload_validation import save_upload

    data = word_bytes(tmp_path)
    destination = tmp_path / "saved.docx"

    assert save_upload(io.BytesIO(data), destination, ".docx", 1_000_000) == len(data)
    assert destination.exists()


@pytest.mark.parametrize("kind", ["junk", "other-zip"])
def test_a_file_that_only_looks_like_word_is_refused_on_upload(tmp_path, kind):
    import io

    from app.core.exceptions import InvalidUploadError
    from app.services.upload_validation import save_upload

    if kind == "junk":
        data = b"PK\x03\x04" + b"x" * 100
    else:
        buffer = io.BytesIO()

        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("notes.txt", "hello")

        data = buffer.getvalue()

    destination = tmp_path / "saved.docx"

    with pytest.raises(InvalidUploadError):
        save_upload(io.BytesIO(data), destination, ".docx", 1_000_000)

    assert not destination.exists()


def test_the_upload_endpoint_accepts_a_word_cv(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from fastapi.testclient import TestClient

    from app.api.dependencies import get_resume_service
    from app.api.routes.v1 import resume_routes
    from app.main import app

    created = []

    class Service:
        def create_resume(self, file_path, filename):
            created.append(filename)

            return SimpleNamespace(
                id="665f1c2b8e4a4b0012345678",
                filename=filename,
                upload_status="pending",
                created_at="2026-01-01T00:00:00",
            )

        def process_resume(self, resume_id):
            pass

    monkeypatch.setattr(resume_routes, "UPLOAD_DIR", tmp_path)
    app.dependency_overrides[get_resume_service] = lambda: Service()

    try:
        response = TestClient(app, raise_server_exceptions=False).post(
            "/api/v1/resumes/upload",
            files={"file": ("cv.docx", word_bytes(tmp_path), "application/octet-stream")},
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 202
    assert created == ["cv.docx"]


def test_the_file_type_is_not_case_sensitive(simple, tmp_path):
    upper = tmp_path / "CV.DOCX"
    upper.write_bytes(simple.read_bytes())

    assert "Backend Engineer" in DocumentParser().parse(upper)["text"]


# ------------------------------------------- a two-column Word template


def docx_two_columns(document, directory: Path):
    """The sidebar layout most Word templates use: one table row, two cells."""

    word = Document()
    sections: dict = {}

    for line in document.truth:
        sections.setdefault(line.section, []).append(line.text)

    table = word.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    first = {"left": True, "right": True}

    def add(cell, side, kind, text):
        if first[side]:
            paragraph = cell.paragraphs[0]
            paragraph.text = ""
            first[side] = False
        else:
            paragraph = cell.add_paragraph()

        if kind == "bullet":
            paragraph.style = "List Bullet"

        if kind == "heading":
            paragraph.style = "Heading 1"

        paragraph.add_run(text)

    for line in sections["personal_info"]:
        add(left, "left", "text", line)

    for section in ("skills", "certifications"):
        add(left, "left", "heading", SECTION_TITLES["en"][section])

        for text in sections[section]:
            add(left, "left", "text", text)

    for section in ("summary", "experience", "education", "projects"):
        add(right, "right", "heading", SECTION_TITLES["en"][section])

        for text in sections[section]:
            if text.startswith("- "):
                add(right, "right", "bullet", text[2:])
            else:
                add(right, "right", "text", text)

    path = directory / f"{document.id}-columns.docx"
    word.save(str(path))

    shown = replace(
        document,
        truth=[
            TruthLine(line.section, "• " + line.text[2:] if line.text.startswith("- ") else line.text)
            for line in document.truth
        ],
    )

    return path, shown


def test_sections_are_found_in_a_two_column_word_template(tmp_path):
    corpus = CVCorpusBuilder(seed=4, per_variant=1).build(tmp_path / "pdf")
    documents = [d for d in corpus if d.wording == "standard" and d.language == "en"][:4]
    results = []

    for document in documents:
        docx, shown = docx_two_columns(document, tmp_path)
        results.append(score_document(shown, pipeline().detect_sections(docx)))

    assert aggregate(results)["all"].f1 >= 0.9


def test_an_external_entity_in_the_xml_is_not_resolved(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("TOP-SECRET-VALUE", encoding="utf-8")

    template = Document()
    path = save(template, tmp_path, "template.docx")
    evil = tmp_path / "evil.docx"

    with zipfile.ZipFile(path) as source, zipfile.ZipFile(evil, "w") as target:
        for name in source.namelist():
            data = source.read(name)

            if name == "word/document.xml":
                text = data.decode("utf-8")
                prolog_end = text.index("?>") + 2
                doctype = (
                    f'\n<!DOCTYPE d [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>'
                )
                body = text[prolog_end:].replace("<w:body>", "<w:body><w:p><w:r><w:t>&xxe;</w:t></w:r></w:p>", 1)
                data = (text[:prolog_end] + doctype + body).encode("utf-8")

            target.writestr(name, data)

    try:
        text = DocxParser.parse(evil)["text"]
    except UnreadableFileError:
        text = ""

    assert "TOP-SECRET-VALUE" not in text
